#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
check_competency_anomalies.py - flag GA7 competency-decision anomalies.

Scans all rubric PDFs in a folder (and its subfolders) and flags, by comparing
the Part B core indicators (C1-C4 ticks) against the overall Competent / Not Yet
Competent decision and the Part A mark (pass = 25/50 = 50%):

  A  <= 2 core ticks but declared Competent
        (decision rule needs at least 3 of 4 -> this contradicts it)
  B  >= 3 core ticks and mark >= 50% but declared Not Yet Competent
        (indicators and mark both support Competent -> likely should be Competent)
  C  >= 3 core ticks and declared Competent but mark < 50%
        (REVIEW only: the rubric allows Competent below 50%, since competency is
         judged independently of the mark - check these are intentional)

Outputs:
  * competency_anomalies.xlsx    - Summary + Anomalies sheets
  * Anomalies/<FLAG>__<file>.pdf - a copy of every flagged rubric, grouped by flag

Runs offline (POPIA-safe). Dependencies: pypdf, openpyxl.
    pip install pypdf openpyxl

Usage:
    python check_competency_anomalies.py <folder>
"""

import argparse
import os
import re
import shutil
import sys

try:
    from pypdf import PdfReader
except ImportError:
    sys.exit("Missing dependency. Run:  pip install pypdf openpyxl")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

C_YES = ["f_c1y", "f_c2y", "f_c3y", "f_c4y"]
C_NO = ["f_c1n", "f_c2n", "f_c3n", "f_c4n"]
OVERALL_PAIR = ("f_overall_comp", "f_overall_notyet")
RUBRIC_MARKERS = set(C_YES + C_NO + ["f_mtotal", "f_student_no", "f_m1", "f_core_count"])
TOTAL_MAX = 50.0
PASS_MARK = 25.0  # 50%
ANOM_DIR = "Anomalies"

HDR = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="005BAA")
BAD_FILL = PatternFill("solid", fgColor="F8D0D5")
WARN_FILL = PatternFill("solid", fgColor="FCE8C8")
GREY_FILL = PatternFill("solid", fgColor="E7E7E7")
FLAG_FILL = {"A": BAD_FILL, "B": BAD_FILL, "C": WARN_FILL, "INCOMPLETE": GREY_FILL}


def norm_sn(raw):
    if raw is None:
        return ""
    s = str(raw).strip()
    if s.endswith(".0"):
        s = s[:-2]
    s = re.sub(r"^[uU]", "", s)
    return re.sub(r"\D", "", s)


def is_checked(v):
    return v not in (None, "", "/Off", "Off", False)


def to_float(v):
    if v in (None, ""):
        return None
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def read(path):
    try:
        fields = PdfReader(path).get_fields() or {}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": str(e)}
    if not (set(fields.keys()) & RUBRIC_MARKERS):
        return {"ok": False, "error": "not-a-rubric"}

    def tx(n):
        f = fields.get(n)
        return "" if f is None or f.get("/V") is None else str(f.get("/V")).strip()

    def btn(n):
        f = fields.get(n)
        return is_checked(f.get("/V")) if f is not None else False

    sn = (norm_sn(tx("f_student_no")) or norm_sn(tx("f_student_no_b"))
          or norm_sn(os.path.splitext(os.path.basename(path))[0]))
    core_ticks = sum(1 for y in C_YES if btn(y))
    comp, nyc = btn(OVERALL_PAIR[0]), btn(OVERALL_PAIR[1])
    if comp and not nyc:
        overall = "Competent"
    elif nyc and not comp:
        overall = "Not Yet Competent"
    else:
        overall = ""
    return {"ok": True, "sn": sn, "core_ticks": core_ticks,
            "overall": overall, "total": to_float(tx("f_mtotal"))}


def scan(folder):
    out = []
    for root, dirs, files in os.walk(folder):
        dirs[:] = [d for d in dirs if d != ANOM_DIR]
        for f in sorted(files):
            if f.lower().endswith(".pdf") and not f.startswith("~"):
                out.append(os.path.join(root, f))
    return out


def classify(ticks, overall, total):
    pct = round(100.0 * total / TOTAL_MAX, 1) if total is not None else None
    if overall == "":
        return "INCOMPLETE", "overall outcome blank or both boxes ticked"
    if ticks <= 2 and overall == "Competent":
        return "A", f"{ticks} core ticks but declared Competent (needs >= 3)"
    if ticks >= 3 and overall == "Not Yet Competent" and total is not None and total >= PASS_MARK:
        return "B", f"{ticks} core ticks and {pct}% but declared Not Yet Competent"
    if ticks >= 3 and overall == "Competent" and total is not None and total < PASS_MARK:
        return "C", f"{ticks} core ticks, Competent, but {pct}% (< 50%) - review (allowed by rubric)"
    if total is None and overall in ("Competent", "Not Yet Competent"):
        return "INCOMPLETE", "total mark missing - mark-based checks skipped"
    return None, None


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def main():
    ap = argparse.ArgumentParser(description="Flag GA7 competency-decision anomalies.")
    ap.add_argument("folder", help="Folder of rubric PDFs (scanned recursively).")
    args = ap.parse_args()

    folder = os.path.abspath(args.folder)
    if not os.path.isdir(folder):
        sys.exit(f"Folder not found: {folder}")

    pdfs = scan(folder)
    if not pdfs:
        sys.exit("No PDF files found.")

    rows = []
    scanned = non_rubric = 0
    counts = {"A": 0, "B": 0, "C": 0, "INCOMPLETE": 0}

    for p in pdfs:
        r = read(p)
        if not r["ok"]:
            non_rubric += 1
            continue
        scanned += 1
        flag, note = classify(r["core_ticks"], r["overall"], r["total"])
        if flag is None:
            continue
        counts[flag] += 1
        total = r["total"]
        rows.append({
            "flag": flag,
            "student_no": ("u" + r["sn"]) if r["sn"] else os.path.basename(p),
            "core_ticks": r["core_ticks"],
            "overall": r["overall"],
            "total": "" if total is None else total,
            "pct": round(100.0 * total / TOTAL_MAX, 1) if total is not None else "",
            "note": note,
            "file": os.path.relpath(p, folder),
            "src": p,
        })

    rows.sort(key=lambda x: (x["flag"], x["student_no"]))

    anom_dir = os.path.join(folder, ANOM_DIR)
    if os.path.isdir(anom_dir):
        shutil.rmtree(anom_dir)
    os.makedirs(anom_dir, exist_ok=True)
    for r in rows:
        dest_name = f"{r['flag']}__{os.path.basename(r['src'])}"
        dest = os.path.join(anom_dir, dest_name)
        n = 1
        while os.path.exists(dest):
            dest = os.path.join(anom_dir, dest_name.replace(".pdf", f"__{n}.pdf"))
            n += 1
        shutil.copy2(r["src"], dest)
        r["copied_as"] = os.path.join(ANOM_DIR, os.path.basename(dest))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sm = wb.create_sheet("Summary")
    sm.append(["GA7 competency-decision anomalies"])
    sm.cell(row=1, column=1).font = Font(bold=True, size=13, color="005BAA")
    sm.append([])
    sm.append(["Rubrics scanned", scanned])
    sm.append(["Non-rubric files ignored", non_rubric])
    sm.append([])
    sm.append(["Flag", "Meaning", "Count"])
    style_header(sm, 3, row=6)
    sm.append(["A", "<= 2 core ticks but declared Competent (contradicts >= 3 rule)", counts["A"]])
    sm.append(["B", ">= 3 core ticks and >= 50% but declared Not Yet Competent", counts["B"]])
    sm.append(["C", ">= 3 core ticks, Competent, but < 50% (review; allowed by rubric)", counts["C"]])
    sm.append(["Incomplete", "outcome not decided or mark missing (not evaluated)", counts["INCOMPLETE"]])
    for i, key in enumerate(["A", "B", "C", "INCOMPLETE"]):
        sm.cell(row=7 + i, column=1).fill = FLAG_FILL[key]
    for col, w in enumerate([12, 62, 8], start=1):
        sm.column_dimensions[get_column_letter(col)].width = w

    ws = wb.create_sheet("Anomalies")
    head = ["Flag", "Student no", "Core ticks", "Overall outcome", "Total/50", "%",
            "Reason", "Source file", "Copied as"]
    ws.append(head)
    for r in rows:
        ws.append([r["flag"], r["student_no"], r["core_ticks"], r["overall"],
                   r["total"], r["pct"], r["note"], r["file"], r["copied_as"]])
        for c in range(1, len(head) + 1):
            ws.cell(row=ws.max_row, column=c).fill = FLAG_FILL[r["flag"]]
    if not rows:
        ws.append(["(no anomalies found)"])
    style_header(ws, len(head))
    for col, w in enumerate([6, 12, 10, 18, 9, 7, 52, 30, 30], start=1):
        ws.column_dimensions[get_column_letter(col)].width = w

    out_xlsx = os.path.join(folder, "competency_anomalies.xlsx")
    wb.save(out_xlsx)

    print(f"Scanned folder : {folder}")
    print(f"  Rubrics scanned  : {scanned}")
    print(f"  Non-rubric files : {non_rubric} (ignored)")
    print(f"  Flag A (<=2 ticks, declared Competent)              : {counts['A']}")
    print(f"  Flag B (>=3 ticks, >=50%, declared Not Yet Competent): {counts['B']}")
    print(f"  Flag C (>=3 ticks, Competent, <50% - review)        : {counts['C']}")
    print(f"  Incomplete (could not fully evaluate)               : {counts['INCOMPLETE']}")
    print(f"\nReport : {out_xlsx}")
    print(f"Copies : {len(rows)} rubric(s) -> {anom_dir}")


if __name__ == "__main__":
    main()
